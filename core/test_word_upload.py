"""Word upload validation and extraction, including legacy .doc routing."""

from io import BytesIO
from pathlib import Path
import subprocess
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from docx import Document
from rest_framework.exceptions import ValidationError

from core.validators import sniff_kind, validate_upload_file
from services.word_extractor import (
    WordExtractionError, WordServiceUnavailable, extract_word_text,
)


OLE_HEADER = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph('Lundi')
    tableau = document.add_table(rows=1, cols=2)
    tableau.cell(0, 0).text = '9 h'
    tableau.cell(0, 1).text = 'Cours de physique'
    document.add_paragraph('Mardi')
    tampon = BytesIO()
    document.save(tampon)
    return tampon.getvalue()


class WordValidationTests(SimpleTestCase):
    def test_docx_valide_est_accepte_sans_consommation_du_flux(self):
        fichier = SimpleUploadedFile('horaire.docx', docx_bytes())
        self.assertEqual(validate_upload_file(fichier), 'docx')
        self.assertEqual(fichier.read(4), b'PK\x03\x04')
        fichier.seek(0)
        self.assertEqual(sniff_kind(fichier.read(16)), 'docx')

    def test_doc_ole_est_accepte(self):
        fichier = SimpleUploadedFile('horaire.doc', OLE_HEADER + b'contenu')
        self.assertEqual(validate_upload_file(fichier), 'doc')
        self.assertEqual(sniff_kind(OLE_HEADER), 'doc')

    def test_zip_sans_document_word_est_refuse(self):
        tampon = BytesIO()
        with ZipFile(tampon, 'w') as archive:
            archive.writestr('autre.txt', 'contenu')
        fichier = SimpleUploadedFile('faux.docx', tampon.getvalue())
        with self.assertRaises(ValidationError):
            validate_upload_file(fichier)

    def test_extension_doc_sur_un_pdf_est_refusee(self):
        fichier = SimpleUploadedFile('faux.doc', b'%PDF-1.4')
        with self.assertRaises(ValidationError):
            validate_upload_file(fichier)


class WordExtractionTests(SimpleTestCase):
    def test_docx_conserve_ordre_des_paragraphes_et_tableau(self):
        with TemporaryDirectory() as dossier:
            chemin = Path(dossier) / 'horaire.docx'
            chemin.write_bytes(docx_bytes())
            texte = extract_word_text(str(chemin), 'docx')
        self.assertLess(texte.index('Lundi'), texte.index('9 h | Cours de physique'))
        self.assertLess(texte.index('9 h | Cours de physique'), texte.index('Mardi'))

    @patch('services.word_extractor.subprocess.run')
    def test_doc_legacy_utilise_antiword(self, executer):
        executer.return_value = subprocess.CompletedProcess([], 0, b'Lundi 9 h physique\n', b'')
        self.assertEqual(extract_word_text('horaire.doc', 'doc'), 'Lundi 9 h physique')
        self.assertEqual(executer.call_args.args[0], ['antiword', '-w', '0', 'horaire.doc'])

    @patch('services.word_extractor.subprocess.run', side_effect=FileNotFoundError)
    def test_convertisseur_absent_signale(self, _executer):
        with self.assertRaises(WordServiceUnavailable):
            extract_word_text('horaire.doc', 'doc')

    @patch('services.word_extractor.subprocess.run')
    def test_doc_illisible_refuse(self, executer):
        executer.return_value = subprocess.CompletedProcess([], 1, b'', b'erreur interne')
        with self.assertRaises(WordExtractionError):
            extract_word_text('horaire.doc', 'doc')
